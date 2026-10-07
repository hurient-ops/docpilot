"""
Phase 0 PoC: 5대 포맷 파서 파이프라인 검증 스크립트
지원 포맷: .docx, .pptx, .xlsx, .pdf, .hwpx
"""

import os

def parse_docx(file_path):
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        return {
            "format": "docx",
            "text": "\n".join(full_text),
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables)
        }
    except Exception as e:
        return {"format": "docx", "text": "", "error": str(e)}

def parse_pptx(file_path):
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        full_text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    full_text.append(shape.text)
        return {
            "format": "pptx",
            "text": "\n".join(full_text),
            "slides": len(prs.slides)
        }
    except Exception as e:
        return {"format": "pptx", "text": "", "error": str(e)}

def parse_xlsx(file_path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        full_text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = [str(cell) for cell in row if cell is not None]
                if row_text:
                    full_text.append(" ".join(row_text))
        return {
            "format": "xlsx",
            "text": "\n".join(full_text),
            "sheets": wb.sheetnames
        }
    except Exception as e:
        return {"format": "xlsx", "text": "", "error": str(e)}

def parse_pdf(file_path):
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        full_text = []
        is_scanned = False
        
        for page_num, page in enumerate(doc):
            text = page.get_text()
            if not text.strip():
                # 텍스트가 없는 스캔본으로 간주하고 OCR 파이프라인 가동
                is_scanned = True
                try:
                    import pytesseract
                    from PIL import Image
                    import io
                    
                    # 렌더링 배율 설정 (해상도 증가)
                    zoom_matrix = fitz.Matrix(2.0, 2.0)
                    pix = page.get_pixmap(matrix=zoom_matrix)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    
                    # OCR 텍스트 추출 (한국어+영어)
                    ocr_text = pytesseract.image_to_string(img, lang="kor+eng")
                    full_text.append(ocr_text)
                    print(f"[Parser] PDF {page_num}페이지 OCR 추출 완료")
                except Exception as e:
                    print(f"[Parser] OCR 실패: {e}")
            else:
                full_text.append(text)
                
        return {
            "format": "pdf",
            "text": "\n".join(full_text),
            "pages": len(doc),
            "ocr_applied": is_scanned
        }
    except Exception as e:
        return {"format": "pdf", "text": "", "error": str(e)}

def parse_hwpx(file_path):
    # HWPX는 OCF(ZIP) 구조의 XML 파일입니다. (임시로 텍스트 추출 시뮬레이션)
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        
        full_text = []
        with zipfile.ZipFile(file_path, 'r') as zf:
            if 'Contents/section0.xml' in zf.namelist():
                xml_content = zf.read('Contents/section0.xml')
                root = ET.fromstring(xml_content)
                # hwp xml 태그에서 텍스트 노드 추출 (네임스페이스 주의)
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        full_text.append(elem.text.strip())
        return {"format": "hwpx", "text": "\n".join(full_text), "status": "parsed"}
    except Exception as e:
        return {"format": "hwpx", "text": "", "error": str(e)}

def generate_preview_pdf_libreoffice(file_path, output_dir):
    """
    LibreOffice Headless 모드를 사용하여 오피스 문서를 PDF로 강제 렌더링.
    (Split-Pane 미리보기 뷰 및 페이지 좌표 계산용)
    """
    import subprocess
    import shutil
    import platform

    soffice = "soffice"
    if platform.system() == "Windows":
        soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"

    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    try:
        # soffice --headless --convert-to pdf --outdir <output_dir> <file_path>
        cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", output_dir, file_path]
        print(f"[Parser] LibreOffice 렌더링 시작: {file_path}")
        subprocess.run(cmd, check=True, timeout=60, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        pdf_path = os.path.join(output_dir, f"{base_name}.pdf")
        if os.path.exists(pdf_path):
            print(f"[Parser] 렌더링 완료: {pdf_path}")
            return pdf_path
        return None
    except Exception as e:
        print(f"[Parser] LibreOffice 렌더링 실패: {e}")
        return None

def parse_document(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.docx':
        return parse_docx(file_path)
    elif ext == '.pptx':
        return parse_pptx(file_path)
    elif ext == '.xlsx':
        return parse_xlsx(file_path)
    elif ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.hwpx':
        return parse_hwpx(file_path)
    else:
        raise ValueError(f"Unsupported format: {ext}")

if __name__ == "__main__":
    print("DocPilot Phase 0 - Parser PoC Initialized.")
    print("Ready to parse 5 major formats.")
